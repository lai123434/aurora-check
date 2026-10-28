"""把 10月30日初稿改成 Rev8（补回字数至 800+，同时保持低 AI 率）。

正文文本严格取自 /workspace/Precheck_Oct30_Rev8_body.txt，
保证生成的 docx 与实测文本逐字一致。
"""
import docx

SRC = "/workspace/BUS100_A1_Annotated_Bibliography_Precheck_Oct30.docx"
DST = "/workspace/Precheck_Oct30_Rev8.docx"
BODY = "/workspace/Precheck_Oct30_Rev8_body.txt"


def main():
    with open(BODY, encoding="utf-8") as f:
        paras = [p.strip() for p in f.read().split("\n\n") if p.strip()]
    assert len(paras) == 9, f"期望 9 个子段，实际 {len(paras)}"

    ANN1 = " ".join(paras[0:3])
    ANN2 = " ".join(paras[3:5])
    ANN3 = " ".join(paras[5:8])
    CONC = paras[8]

    # 段落索引对齐原 docx：5/7/9 = 三条注释，10 = 总结（4/6/8 是 APA 引用条目，不动）
    targets = {5: ANN1, 7: ANN2, 9: ANN3, 10: CONC}

    d = docx.Document(SRC)
    replaced = 0
    for i, p in enumerate(d.paragraphs):
        if i in targets:
            new = targets[i]
            if p.runs:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run(new)
            replaced += 1
    d.save(DST)

    # 校验最终 docx 词数（去封皮）
    total = sum(len(p.text.split()) for p in d.paragraphs)
    cover = sum(len(p.text.split()) for p in d.paragraphs[:4])
    print(f"已生成 {DST}，替换正文段落 {replaced}/{len(targets)} 处")
    print(f"总词数 {total}，去封皮 {total - cover} 词（目标 800+）")


if __name__ == "__main__":
    main()
