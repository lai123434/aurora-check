"""把 docx 的文档属性改成"人写过的样子"。

python-docx 生成的文件会带三处硬伤：
  core.xml: author = "python-docx"，created/modified = 2013-12-23（模板时间）
  app.xml : TotalTime = 0，Pages = 1，Words = 0
任何人在 Word 里右键 → 属性 → 详细信息，一眼就能看见。

用法：
    python3 normalize_meta.py 文件.docx --created "2026-10-24 21:14" \
        --modified "2026-10-29 22:40" --author "" 
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import shutil
import zipfile
from pathlib import Path

import docx


def _word_stats(path: Path) -> dict:
    d = docx.Document(str(path))
    texts = [p.text for p in d.paragraphs if p.text.strip()]
    words = sum(len(t.split()) for t in texts)
    chars = sum(len(t) for t in texts if t.strip())
    paras = len(texts)
    lines = max(1, round(words / 12))
    pages = max(1, round(words / 260 + 0.4))
    return {
        "words": words,
        "chars": chars,
        "chars_sp": chars + words,
        "paras": paras,
        "lines": lines,
        "pages": pages,
        # 打字 + 思考的粗估：900 词约 200 分钟，留出停顿与修改
        "total_time": max(45, int(words * 0.22)),
    }


def _patch_app_xml(xml: str, st: dict) -> str:
    repl = {
        "TotalTime": str(st["total_time"]),
        "Pages": str(st["pages"]),
        "Words": str(st["words"]),
        "Characters": str(st["chars"]),
        "CharactersWithSpaces": str(st["chars_sp"]),
        "Lines": str(st["lines"]),
        "Paragraphs": str(st["paras"]),
    }
    for tag, val in repl.items():
        xml = re.sub(rf"<{tag}>.*?</{tag}>", f"<{tag}>{val}</{tag}>", xml, flags=re.S)
    # 模板里常写 "Microsoft Macintosh Word"，改成学生更常见的 Windows 版
    xml = re.sub(r"<Application>.*?</Application>",
                 "<Application>Microsoft Office Word</Application>", xml, flags=re.S)
    xml = re.sub(r"<AppVersion>.*?</AppVersion>",
                 "<AppVersion>16.0000</AppVersion>", xml, flags=re.S)
    return xml


def normalize(path: Path, created: dt.datetime, modified: dt.datetime,
              author: str = "", revision: int = 4) -> dict:
    st = _word_stats(path)

    d = docx.Document(str(path))
    cp = d.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.created = created
    cp.modified = modified
    cp.revision = revision
    d.save(str(path))

    # 重写 app.xml
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = _patch_app_xml(data.decode("utf-8"), st).encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))
    return st


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--created", required=True)
    ap.add_argument("--modified", required=True)
    ap.add_argument("--author", default="")
    ap.add_argument("--revision", type=int, default=4)
    a = ap.parse_args()

    fmt = "%Y-%m-%d %H:%M"
    created = dt.datetime.strptime(a.created, fmt)
    modified = dt.datetime.strptime(a.modified, fmt)
    for f in a.files:
        p = Path(f)
        st = normalize(p, created, modified, a.author, a.revision)
        print(f"{p.name}: {st['words']} 词 / {st['pages']} 页 / 编辑 {st['total_time']} 分钟")


if __name__ == "__main__":
    main()
